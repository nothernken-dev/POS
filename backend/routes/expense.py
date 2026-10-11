from flask import request, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models.expense import Expense
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import qn, OxmlElement


def get_current_shop_id():
    if hasattr(current_user, "id"):
        return current_user.id
    return None


def init_expense_routes(app):

    @app.route("/api/expenses", methods=["GET"])
    @login_required
    def get_expenses():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            status = request.args.get("status")
            payment_method = request.args.get("payment_method")
            start_date = request.args.get("start_date")
            end_date = request.args.get("end_date")
            search = request.args.get("search")

            query = Expense.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            if payment_method:
                query = query.filter_by(payment_method=payment_method)
            if start_date:
                query = query.filter(Expense.date >= start_date)
            if end_date:
                query = query.filter(Expense.date <= end_date)
            if search:
                query = query.filter(Expense.item_name.contains(search))

            expenses = query.order_by(Expense.date.desc()).all()
            return jsonify([e.to_dict() for e in expenses])
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses/<int:expense_id>", methods=["GET"])
    @login_required
    def get_expense(expense_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            expense = Expense.query.filter_by(id=expense_id, shop_id=shop_id).first()
            if not expense:
                return jsonify({"error": "Expense not found"}), 404
            return jsonify(expense.to_dict())
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses", methods=["POST"])
    @login_required
    def create_expense():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            data = request.get_json()

            required = ["item_name", "amount", "date", "payment_method"]
            missing = [f for f in required if f not in data]
            if missing:
                return jsonify({"error": f'Missing required fields: {", ".join(missing)}'}), 400

            try:
                amount = float(data["amount"])
                if amount < 0:
                    return jsonify({"error": "Amount cannot be negative"}), 400
            except ValueError:
                return jsonify({"error": "Invalid amount format"}), 400

            expense = Expense(
                shop_id=shop_id,
                item_name=data["item_name"],
                amount=amount,
                date=datetime.strptime(data["date"], "%Y-%m-%d").date(),
                payment_method=data["payment_method"],
                reference=data.get("reference", ""),
                status=data.get("status", "Pending"),
                notes=data.get("notes", ""),
                created_by=current_user.id if current_user.is_authenticated else None,
            )
            db.session.add(expense)
            db.session.commit()
            return jsonify(expense.to_dict()), 201
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses/<int:expense_id>", methods=["PUT"])
    @login_required
    def update_expense(expense_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            expense = Expense.query.filter_by(id=expense_id, shop_id=shop_id).first()
            if not expense:
                return jsonify({"error": "Expense not found"}), 404

            data = request.get_json()
            if "item_name" in data:
                expense.item_name = data["item_name"]
            if "amount" in data:
                amount = float(data["amount"])
                if amount < 0:
                    return jsonify({"error": "Amount cannot be negative"}), 400
                expense.amount = amount
            if "date" in data:
                expense.date = datetime.strptime(data["date"], "%Y-%m-%d").date()
            if "payment_method" in data:
                expense.payment_method = data["payment_method"]
            if "reference" in data:
                expense.reference = data["reference"]
            if "status" in data:
                expense.status = data["status"]
            if "notes" in data:
                expense.notes = data["notes"]

            db.session.commit()
            return jsonify(expense.to_dict())
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses/<int:expense_id>", methods=["DELETE"])
    @login_required
    def delete_expense(expense_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            expense = Expense.query.filter_by(id=expense_id, shop_id=shop_id).first()
            if not expense:
                return jsonify({"error": "Expense not found"}), 404
            db.session.delete(expense)
            db.session.commit()
            return jsonify({"message": "Expense deleted successfully"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses/stats", methods=["GET"])
    @login_required
    def get_expense_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            start_date = request.args.get("start_date")
            end_date = request.args.get("end_date")

            query = Expense.query.filter_by(shop_id=shop_id)
            if start_date:
                query = query.filter(Expense.date >= start_date)
            if end_date:
                query = query.filter(Expense.date <= end_date)
            expenses = query.all()

            total = sum(float(e.amount or 0) for e in expenses)
            paid = sum(float(e.amount or 0) for e in expenses if e.status == "Paid")
            pending = sum(float(e.amount or 0) for e in expenses if e.status == "Pending")
            overdue = sum(float(e.amount or 0) for e in expenses if e.status == "Overdue")

            return jsonify({
                "total": total, "paid": paid, "pending": pending, "overdue": overdue,
                "count": len(expenses),
                "average": total / len(expenses) if expenses else 0,
            })
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/expenses/export/<string:format>", methods=["GET"])
    @login_required
    def export_expenses(format):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            status = request.args.get("status")
            payment_method = request.args.get("payment_method")
            start_date = request.args.get("start_date")
            end_date = request.args.get("end_date")
            search = request.args.get("search")

            query = Expense.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            if payment_method:
                query = query.filter_by(payment_method=payment_method)
            if start_date:
                query = query.filter(Expense.date >= start_date)
            if end_date:
                query = query.filter(Expense.date <= end_date)
            if search:
                query = query.filter(Expense.item_name.contains(search))

            expenses = query.order_by(Expense.date.desc()).all()
            if not expenses:
                return jsonify({"error": "No expenses data to export."}), 404

            if format == "excel":
                return export_excel(expenses)
            elif format == "pdf":
                return export_pdf(expenses)
            elif format == "word":
                return export_word(expenses)
            return jsonify({"error": "Invalid format"}), 400
        except Exception as e:
            return jsonify({"error": f"Failed to export: {str(e)}"}), 500

    def export_excel(expenses):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Expenses Report"
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(horizontal="left", vertical="center")
        number_alignment = Alignment(horizontal="right", vertical="center")
        border = Border(left=Side(style="thin"), right=Side(style="thin"),
                        top=Side(style="thin"), bottom=Side(style="thin"))

        title_cell = ws.cell(row=1, column=1, value="EXPENSES REPORT")
        title_cell.font = Font(bold=True, size=16)
        ws.merge_cells("A1:I1")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        total_expenses = sum(float(e.amount or 0) for e in expenses)
        paid = sum(float(e.amount or 0) for e in expenses if e.status == "Paid")
        pending = sum(float(e.amount or 0) for e in expenses if e.status == "Pending")
        overdue = sum(float(e.amount or 0) for e in expenses if e.status == "Overdue")

        ws.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        ws.cell(row=3, column=1, value=f"Total: {len(expenses)}")
        ws.cell(row=3, column=3, value=f"Amount: KES {total_expenses:,.2f}")
        ws.cell(row=3, column=5, value=f"Paid: KES {paid:,.2f}")
        ws.cell(row=3, column=7, value=f"Pending: KES {pending:,.2f}")
        ws.cell(row=3, column=9, value=f"Overdue: KES {overdue:,.2f}")

        headers = ["ID", "Item Name", "Amount (KES)", "Date", "Payment Method",
                   "Reference", "Status", "Notes", "Created At"]
        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        for row_idx, expense in enumerate(expenses, header_row + 1):
            row_data = [
                expense.id, expense.item_name, float(expense.amount or 0),
                expense.date.strftime("%Y-%m-%d") if expense.date else "N/A",
                expense.payment_method, expense.reference or "-",
                expense.status, expense.notes or "-",
                expense.created_at.strftime("%Y-%m-%d %H:%M") if expense.created_at else "N/A",
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = border
                if isinstance(value, (int, float)):
                    cell.alignment = number_alignment
                    if col == 3:
                        cell.number_format = "#,##0.00"
                else:
                    cell.alignment = cell_alignment

        for col in ws.columns:
            max_length = 0
            column = col[0].column
            for cell in col:
                try:
                    if isinstance(cell, openpyxl.cell.cell.MergedCell):
                        continue
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 12), 30)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True,
                         download_name=f'expenses_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')

    def export_pdf(expenses):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter),
                                rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("T", parent=styles["Heading1"], fontSize=18, alignment=1,
                                     spaceAfter=10, textColor=colors.HexColor("#8B0000"))
        elements = [Paragraph("Expenses Report", title_style)]
        date_style = ParagraphStyle("D", parent=styles["Normal"], fontSize=10, alignment=1, spaceAfter=15)
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", date_style))

        total_expenses = sum(float(e.amount or 0) for e in expenses)
        summary_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10, spaceAfter=15)
        elements.append(Paragraph(f"<b>Total:</b> {len(expenses)} | <b>Amount:</b> KES {total_expenses:,.2f}", summary_style))

        table_data = [["ID", "Item", "Amount", "Date", "Payment", "Ref", "Status"]]
        for e in expenses:
            table_data.append([
                str(e.id), e.item_name[:25], f"KES {float(e.amount or 0):,.2f}",
                e.date.strftime("%Y-%m-%d") if e.date else "N/A",
                e.payment_method, e.reference or "-", e.status,
            ])

        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8B0000")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#cccccc")),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return send_file(buffer, mimetype="application/pdf", as_attachment=True,
                         download_name=f'expenses_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf')

    def export_word(expenses):
        doc = docx.Document()
        title = doc.add_heading("Expenses Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        total_expenses = sum(float(e.amount or 0) for e in expenses)
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph(f"Total: {len(expenses)} | Amount: KES {total_expenses:,.2f}")

        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        for i, h in enumerate(["ID", "Item", "Amount", "Date", "Payment", "Ref", "Status"]):
            table.rows[0].cells[i].text = h

        for e in expenses:
            row = table.add_row().cells
            row[0].text = str(e.id)
            row[1].text = e.item_name
            row[2].text = f"KES {float(e.amount or 0):,.2f}"
            row[3].text = e.date.strftime("%Y-%m-%d") if e.date else "N/A"
            row[4].text = e.payment_method
            row[5].text = e.reference or "-"
            row[6].text = e.status

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return send_file(buffer,
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         as_attachment=True,
                         download_name=f'expenses_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx')