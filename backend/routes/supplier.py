from flask import request, jsonify, current_app
from flask_login import login_required, current_user
from extensions import db
from models.supplier import Supplier
from models.inventory import Product
from models.purchase_order import PurchaseOrder, PurchaseOrderItem, OrderStatusHistory
from models.shop import Shop
from datetime import datetime, timezone
import logging

logger = logging.getLogger(__name__)


def utcnow():
    return datetime.now(timezone.utc)


def get_current_shop_id():
    if hasattr(current_user, "id"):
        return current_user.id
    return None


def get_shop_name(shop_id):
    if not shop_id:
        return "System"
    shop = Shop.query.filter_by(id=shop_id).first()
    return shop.name if shop else f"Shop #{shop_id}"


def get_available_status_transitions(current_status):
    transitions = {
        "Pending": ["Ordered", "Cancelled"],
        "Ordered": ["Received", "Cancelled", "Pending"],
        "Received": ["Pending", "Ordered"],
        "Cancelled": ["Pending", "Ordered"],
    }
    return transitions.get(current_status, [])


def validate_status_transition(current_status, new_status):
    return new_status in get_available_status_transitions(current_status)


def enrich_order_with_shop_name(order_dict, order):
    order_dict["created_by_name"] = get_shop_name(order.created_by)
    return order_dict


def enrich_history_with_shop_name(history_dict, history):
    history_dict["changed_by_name"] = get_shop_name(history.changed_by)
    return history_dict


def init_supplier_routes(app):

    @app.route("/api/suppliers", methods=["GET"])
    @login_required
    def get_suppliers():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            status = request.args.get("status")
            item_name = request.args.get("item_name")
            search = request.args.get("search")

            query = Supplier.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            if item_name:
                query = query.filter_by(item_name=item_name)
            if search:
                query = query.filter(db.or_(
                    Supplier.name.contains(search),
                    Supplier.contact_person.contains(search),
                    Supplier.email.contains(search),
                    Supplier.phone.contains(search),
                ))

            return jsonify([s.to_dict() for s in query.order_by(Supplier.created_at.desc()).all()])
        except Exception as e:
            logger.error(f"Error fetching suppliers: {e}")
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/<int:supplier_id>", methods=["GET"])
    @login_required
    def get_supplier(supplier_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            supplier = Supplier.query.filter_by(id=supplier_id, shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404
            return jsonify(supplier.to_dict())
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers", methods=["POST"])
    @login_required
    def create_supplier():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            data = request.get_json()

            required = ["name", "contact_person", "email", "phone"]
            missing = [f for f in required if f not in data]
            if missing:
                return jsonify({"error": f'Missing required fields: {", ".join(missing)}'}), 400

            if Supplier.query.filter_by(email=data["email"], shop_id=shop_id).first():
                return jsonify({"error": "Supplier with this email already exists"}), 400

            supplier = Supplier(
                shop_id=shop_id, name=data["name"],
                contact_person=data["contact_person"], email=data["email"],
                phone=data["phone"], address=data.get("address", ""),
                city=data.get("city", ""), state=data.get("state", ""),
                country=data.get("country", ""), postal_code=data.get("postal_code", ""),
                item_name=data.get("item_name", ""), status=data.get("status", "Active"),
                rating=float(data.get("rating", 0)), notes=data.get("notes", ""),
            )
            db.session.add(supplier)
            db.session.commit()
            return jsonify(supplier.to_dict()), 201
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/<int:supplier_id>", methods=["PUT"])
    @login_required
    def update_supplier(supplier_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            supplier = Supplier.query.filter_by(id=supplier_id, shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404

            data = request.get_json()
            if "email" in data:
                existing = Supplier.query.filter_by(email=data["email"], shop_id=shop_id).first()
                if existing and existing.id != supplier_id:
                    return jsonify({"error": "Email already in use"}), 400

            for field in ["name", "contact_person", "email", "phone", "address",
                          "city", "state", "country", "postal_code", "item_name",
                          "status", "notes"]:
                if field in data:
                    setattr(supplier, field, data[field])
            if "rating" in data:
                supplier.rating = float(data["rating"])

            db.session.commit()
            return jsonify(supplier.to_dict())
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/<int:supplier_id>", methods=["DELETE"])
    @login_required
    def delete_supplier(supplier_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            supplier = Supplier.query.filter_by(id=supplier_id, shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404
            supplier.status = "Inactive"
            db.session.commit()
            return jsonify({"message": "Supplier deactivated successfully"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/stats", methods=["GET"])
    @login_required
    def get_supplier_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            suppliers = Supplier.query.filter_by(shop_id=shop_id).all()

            category_breakdown = {}
            for s in suppliers:
                if s.item_name:
                    category_breakdown[s.item_name] = category_breakdown.get(s.item_name, 0) + 1

            return jsonify({
                "total": len(suppliers),
                "active": Supplier.query.filter_by(shop_id=shop_id, status="Active").count(),
                "inactive": Supplier.query.filter_by(shop_id=shop_id, status="Inactive").count(),
                "pending": Supplier.query.filter_by(shop_id=shop_id, status="Pending").count(),
                "total_products": sum(s.total_products or 0 for s in suppliers),
                "total_orders": sum(s.total_orders or 0 for s in suppliers),
                "total_spent": sum(float(s.total_spent or 0) for s in suppliers),
                "category_breakdown": category_breakdown,
            })
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/update-stats", methods=["POST"])
    @login_required
    def update_supplier_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            suppliers = Supplier.query.filter_by(shop_id=shop_id).all()
            for s in suppliers:
                s.total_products = Product.query.filter_by(supplier_id=s.id, shop_id=shop_id).count()
            db.session.commit()
            return jsonify({"message": f"Updated {len(suppliers)} suppliers", "updated_count": len(suppliers)}), 200
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/<int:supplier_id>/update-stats", methods=["POST"])
    @login_required
    def update_single_supplier_stats(supplier_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            supplier = Supplier.query.filter_by(id=supplier_id, shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404
            supplier.total_products = Product.query.filter_by(supplier_id=supplier_id, shop_id=shop_id).count()
            db.session.commit()
            return jsonify(supplier.to_dict())
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/<int:supplier_id>/products", methods=["GET"])
    @login_required
    def get_supplier_products(supplier_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            supplier = Supplier.query.filter_by(id=supplier_id, shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404
            products = Product.query.filter_by(supplier_id=supplier_id, shop_id=shop_id).all()
            return jsonify([p.to_dict() for p in products])
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/suppliers/names", methods=["GET"])
    @login_required
    def get_supplier_names():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            suppliers = Supplier.query.filter_by(shop_id=shop_id, status="Active").order_by(Supplier.name).all()
            return jsonify([{"id": s.id, "name": s.name, "item_name": s.item_name} for s in suppliers])
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders", methods=["POST"])
    @login_required
    def create_purchase_order():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            data = request.get_json()
            if not data or "supplier_id" not in data:
                return jsonify({"error": "Supplier ID is required"}), 400
            if "items" not in data or len(data["items"]) == 0:
                return jsonify({"error": "At least one item is required"}), 400

            supplier = Supplier.query.filter_by(id=data["supplier_id"], shop_id=shop_id).first()
            if not supplier:
                return jsonify({"error": "Supplier not found"}), 404

            subtotal = 0
            items = []
            for item_data in data["items"]:
                if "product_id" in item_data and item_data["product_id"]:
                    product = Product.query.filter_by(id=item_data["product_id"], shop_id=shop_id).first()
                    if not product:
                        return jsonify({"error": f'Product {item_data["product_id"]} not found'}), 404
                    price = float(item_data.get("price", product.price))
                    quantity = item_data["quantity"]
                    items.append({
                        "product": product, "price": price, "quantity": quantity,
                        "total": price * quantity,
                        "product_name": product.name, "product_sku": product.sku,
                    })
                else:
                    product_name = item_data.get("product_name", "")
                    if not product_name:
                        return jsonify({"error": "Product name required for manual entry"}), 400
                    product = Product.query.filter_by(name=product_name, shop_id=shop_id).first()
                    price = float(item_data.get("price", 0))
                    quantity = item_data.get("quantity", 1)
                    if product:
                        product_sku = product.sku
                    else:
                        product_sku = f"PO-{utcnow().strftime('%Y%m%d')}-{len(items)+1}"
                    items.append({
                        "product": product, "price": price, "quantity": quantity,
                        "total": price * quantity,
                        "product_name": product_name, "product_sku": product_sku,
                    })
                subtotal += items[-1]["total"]

            tax = float(data.get("tax", subtotal * 0.16))
            discount = float(data.get("discount", 0))
            total = subtotal + tax - discount

            order = PurchaseOrder(
                shop_id=shop_id, supplier_id=data["supplier_id"],
                order_date=datetime.strptime(
                    data.get("order_date", datetime.now().strftime("%Y-%m-%d")), "%Y-%m-%d"
                ).date(),
                subtotal=subtotal, tax=tax, discount=discount, total=total,
                notes=data.get("notes", ""), status="Pending",
                created_by=current_user.id if current_user.is_authenticated else None,
            )
            order.order_number = order.generate_order_number()
            db.session.add(order)
            db.session.flush()

            db.session.add(OrderStatusHistory(
                order_id=order.id, old_status="", new_status="Pending",
                changed_by=current_user.id if current_user.is_authenticated else None,
                notes="Order created",
            ))

            for item_data in items:
                db.session.add(PurchaseOrderItem(
                    shop_id=shop_id, order_id=order.id,
                    product_id=item_data["product"].id if item_data["product"] else None,
                    product_name=item_data["product_name"],
                    product_sku=item_data["product_sku"],
                    quantity=item_data["quantity"], price=item_data["price"],
                    total=item_data["total"],
                ))

            supplier.total_orders = (supplier.total_orders or 0) + 1
            supplier.last_order_date = utcnow()
            db.session.commit()

            return jsonify(enrich_order_with_shop_name(order.to_dict(), order)), 201
        except Exception as e:
            db.session.rollback()
            logger.error(f"Error creating PO: {e}")
            return jsonify({"error": f"Failed to create PO: {str(e)}"}), 500

    @app.route("/api/purchase-orders", methods=["GET"])
    @login_required
    def get_purchase_orders():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            query = PurchaseOrder.query.filter_by(shop_id=shop_id)
            status = request.args.get("status")
            supplier_id = request.args.get("supplier_id")
            if status and status != "All":
                query = query.filter_by(status=status)
            if supplier_id:
                query = query.filter_by(supplier_id=supplier_id)

            orders = query.order_by(PurchaseOrder.created_at.desc()).all()
            return jsonify([enrich_order_with_shop_name(o.to_dict(), o) for o in orders]), 200
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/<int:order_id>", methods=["GET"])
    @login_required
    def get_purchase_order(order_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404
            return jsonify(enrich_order_with_shop_name(order.to_dict(), order)), 200
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/<int:order_id>/status", methods=["PATCH"])
    @login_required
    def update_order_status(order_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404

            data = request.get_json()
            new_status = data.get("status")
            if not new_status:
                return jsonify({"error": "Status is required"}), 400
            if new_status not in ["Pending", "Ordered", "Received", "Cancelled"]:
                return jsonify({"error": "Invalid status"}), 400
            if not validate_status_transition(order.status, new_status):
                return jsonify({"error": f"Cannot transition from {order.status} to {new_status}"}), 400

            old_status = order.status

            if old_status == "Received" and new_status != "Received":
                for item in order.items:
                    product = Product.query.filter_by(id=item.product_id, shop_id=shop_id).first()
                    if product:
                        product.stock = max(0, (product.stock or 0) - item.quantity)

            if new_status == "Received" and old_status != "Received":
                for item in order.items:
                    product = None
                    if item.product_id:
                        product = Product.query.filter_by(id=item.product_id, shop_id=shop_id).first()
                    if not product:
                        product = Product.query.filter_by(name=item.product_name, shop_id=shop_id).first()
                    if not product:
                        sku_parts = item.product_name[:3].upper()
                        product = Product(
                            shop_id=shop_id, name=item.product_name,
                            sku=f"{sku_parts}-{utcnow().strftime('%Y%m%d%H%M%S')}",
                            price=item.price, stock=0,
                            description=f"Added from PO {order.order_number}",
                            supplier_id=order.supplier_id,
                        )
                        db.session.add(product)
                        db.session.flush()
                        item.product_id = product.id
                    product.stock = (product.stock or 0) + item.quantity

            order.status = new_status
            history = OrderStatusHistory(
                order_id=order.id, old_status=old_status, new_status=new_status,
                changed_by=current_user.id if current_user.is_authenticated else None,
                notes=f"Status changed from {old_status} to {new_status}",
            )
            db.session.add(history)
            db.session.commit()

            return jsonify({
                "success": True,
                "message": f"Order status updated to {new_status}",
                "order": enrich_order_with_shop_name(order.to_dict(), order),
                "history": enrich_history_with_shop_name(history.to_dict(), history),
            }), 200
        except Exception as e:
            db.session.rollback()
            logger.error(f"Error updating status: {e}")
            return jsonify({"error": f"Failed to update: {str(e)}"}), 500

    @app.route("/api/purchase-orders/<int:order_id>/history", methods=["GET"])
    @login_required
    def get_order_history(order_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404
            history = OrderStatusHistory.query.filter_by(order_id=order.id) \
                .order_by(OrderStatusHistory.created_at.desc()).all()
            return jsonify([enrich_history_with_shop_name(h.to_dict(), h) for h in history]), 200
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/<int:order_id>", methods=["DELETE"])
    @login_required
    def delete_purchase_order(order_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404

            if order.status == "Received":
                time_diff = utcnow() - order.updated_at if order.updated_at else None
                if time_diff and time_diff.total_seconds() < 3600:
                    return jsonify({"warning": "Order received recently. Consider reverting.", "confirm": True}), 400
                return jsonify({"error": "Cannot cancel received order older than 1 hour"}), 400

            old_status = order.status
            order.status = "Cancelled"
            db.session.add(OrderStatusHistory(
                order_id=order.id, old_status=old_status, new_status="Cancelled",
                changed_by=current_user.id if current_user.is_authenticated else None,
                notes="Order cancelled",
            ))
            db.session.commit()
            return jsonify({"message": "Purchase order cancelled"}), 200
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/stats", methods=["GET"])
    @login_required
    def get_purchase_order_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            all_orders = PurchaseOrder.query.filter_by(shop_id=shop_id).all()
            return jsonify({
                "total": len(all_orders),
                "pending": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Pending").count(),
                "ordered": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Ordered").count(),
                "received": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Received").count(),
                "cancelled": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Cancelled").count(),
                "total_amount": sum(float(o.total or 0) for o in all_orders),
            })
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/status-stats", methods=["GET"])
    @login_required
    def get_order_status_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            recent_changes = OrderStatusHistory.query.join(PurchaseOrder).filter(
                PurchaseOrder.shop_id == shop_id
            ).order_by(OrderStatusHistory.created_at.desc()).limit(10).all()

            return jsonify({
                "total": PurchaseOrder.query.filter_by(shop_id=shop_id).count(),
                "pending": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Pending").count(),
                "ordered": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Ordered").count(),
                "received": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Received").count(),
                "cancelled": PurchaseOrder.query.filter_by(shop_id=shop_id, status="Cancelled").count(),
                "recent_changes": [enrich_history_with_shop_name(h.to_dict(), h) for h in recent_changes],
            }), 200
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/fix-order-histories", methods=["POST"])
    @login_required
    def fix_order_histories():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            orders = PurchaseOrder.query.filter_by(shop_id=shop_id).all()
            fixed = 0
            for order in orders:
                if OrderStatusHistory.query.filter_by(order_id=order.id).count() == 0:
                    db.session.add(OrderStatusHistory(
                        order_id=order.id, old_status="", new_status=order.status,
                        changed_by=current_user.id if current_user.is_authenticated else None,
                        notes="Order created (fix)",
                        created_at=order.created_at or utcnow(),
                    ))
                    fixed += 1
            db.session.commit()
            return jsonify({"success": True, "message": f"Fixed {fixed} orders", "fixed_count": fixed}), 200
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/order-status-transitions", methods=["GET"])
    @login_required
    def get_status_transitions():
        try:
            return jsonify({
                "success": True,
                "transitions": {
                    "Pending": get_available_status_transitions("Pending"),
                    "Ordered": get_available_status_transitions("Ordered"),
                    "Received": get_available_status_transitions("Received"),
                    "Cancelled": get_available_status_transitions("Cancelled"),
                },
            }), 200
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/<int:order_id>/export/pdf", methods=["GET"])
    @login_required
    def export_order_pdf(order_id):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.enums import TA_LEFT, TA_CENTER
            import io

            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404

            shop = Shop.query.filter_by(id=shop_id).first()
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72,
                                    topMargin=72, bottomMargin=72)
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle("T", parent=styles["Heading1"], fontSize=22,
                                         textColor=colors.HexColor("#1a56db"), alignment=TA_LEFT)
            subtitle_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10,
                                            textColor=colors.HexColor("#6b7280"), alignment=TA_LEFT)

            content = []
            content.append(Paragraph(f"<b>{shop.name if shop else 'My Shop'}</b>", title_style))
            if shop and shop.address:
                content.append(Paragraph(shop.address, subtitle_style))
            content.append(Spacer(1, 12))

            info_data = [
                ["PO Number:", order.order_number, "Supplier:", order.supplier.name if order.supplier else "N/A"],
                ["Date:", order.order_date.strftime("%B %d, %Y") if order.order_date else "N/A",
                 "Contact:", order.supplier.contact_person if order.supplier else "N/A"],
                ["Status:", order.status, "Email:", order.supplier.email if order.supplier else "N/A"],
                ["Created By:", get_shop_name(order.created_by), "Phone:", order.supplier.phone if order.supplier else "N/A"],
            ]
            info_table = Table(info_data, colWidths=[80, 160, 80, 160])
            info_table.setStyle(TableStyle([
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#6b7280")),
                ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#6b7280")),
            ]))
            content.append(info_table)
            content.append(Spacer(1, 15))

            item_data = [["#", "Product", "SKU", "Qty", "Price", "Total"]]
            for idx, item in enumerate(order.items, 1):
                item_data.append([
                    str(idx), item.product_name, item.product_sku, str(item.quantity),
                    f"KES {float(item.price or 0):,.2f}",
                    f"KES {float(item.total or 0):,.2f}",
                ])
            items_table = Table(item_data, colWidths=[25, 200, 80, 40, 80, 100])
            items_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
                ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ]))
            content.append(items_table)
            content.append(Spacer(1, 10))

            totals_data = [
                ["Subtotal:", f"KES {float(order.subtotal or 0):,.2f}"],
                ["Tax:", f"KES {float(order.tax or 0):,.2f}"],
                ["Discount:", f"KES {float(order.discount or 0):,.2f}"],
                ["TOTAL:", f"KES {float(order.total or 0):,.2f}"],
            ]
            totals_table = Table(totals_data, colWidths=[100, 100])
            totals_table.setStyle(TableStyle([
                ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, -1), (-1, -1), 12),
                ("TEXTCOLOR", (1, -1), (1, -1), colors.HexColor("#1a56db")),
            ]))
            content.append(totals_table)

            if order.notes:
                content.append(Spacer(1, 15))
                content.append(Paragraph(f"<b>Notes:</b> {order.notes}", styles["Normal"]))

            content.append(Spacer(1, 30))
            content.append(Paragraph("Thank you for your business!",
                                     ParagraphStyle("F", parent=styles["Normal"], fontSize=8,
                                                    alignment=TA_CENTER, textColor=colors.HexColor("#9ca3af"))))

            doc.build(content)
            pdf_data = buffer.getvalue()
            buffer.close()

            response = current_app.response_class(pdf_data, mimetype="application/pdf")
            response.headers["Content-Disposition"] = f"attachment; filename=INVOICE-{order.order_number}.pdf"
            return response
        except Exception as e:
            logger.error(f"PDF export error: {e}")
            return jsonify({"error": str(e)}), 500

    @app.route("/api/purchase-orders/<int:order_id>/export/excel", methods=["GET"])
    @login_required
    def export_order_excel(order_id):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            import io

            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            order = PurchaseOrder.query.filter_by(id=order_id, shop_id=shop_id).first()
            if not order:
                return jsonify({"error": "Purchase order not found"}), 404

            shop = Shop.query.filter_by(id=shop_id).first()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"PO-{order.order_number}"

            title_font = Font(bold=True, size=14, color="1a56db")
            header_font = Font(bold=True, size=10, color="374151")
            header_fill = PatternFill(start_color="f3f4f6", end_color="f3f4f6", fill_type="solid")
            border = Border(left=Side(style="thin"), right=Side(style="thin"),
                            top=Side(style="thin"), bottom=Side(style="thin"))

            ws.merge_cells("A1:F1")
            ws["A1"] = shop.name if shop else "My Shop"
            ws["A1"].font = title_font

            row = 5
            data_pairs = [
                ("PO Number:", order.order_number, "Supplier:", order.supplier.name if order.supplier else "N/A"),
                ("Date:", order.order_date.strftime("%B %d, %Y") if order.order_date else "N/A",
                 "Contact:", order.supplier.contact_person if order.supplier else "N/A"),
                ("Status:", order.status, "Email:", order.supplier.email if order.supplier else "N/A"),
                ("Created By:", get_shop_name(order.created_by),
                 "Phone:", order.supplier.phone if order.supplier else "N/A"),
            ]
            for left, lv, right, rv in data_pairs:
                ws.cell(row=row, column=1, value=left).font = Font(bold=True, color="6b7280", size=8)
                ws.cell(row=row, column=2, value=lv).font = Font(size=8)
                ws.cell(row=row, column=3, value=right).font = Font(bold=True, color="6b7280", size=8)
                ws.cell(row=row, column=4, value=rv).font = Font(size=8)
                row += 1

            row += 2
            headers = ["#", "Product", "SKU", "Quantity", "Price", "Total"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
            row += 1

            for idx, item in enumerate(order.items, 1):
                ws.cell(row=row, column=1, value=idx).border = border
                ws.cell(row=row, column=2, value=item.product_name).border = border
                ws.cell(row=row, column=3, value=item.product_sku).border = border
                ws.cell(row=row, column=4, value=item.quantity).border = border
                ws.cell(row=row, column=5, value=f"KES {float(item.price or 0):,.2f}").border = border
                ws.cell(row=row, column=6, value=f"KES {float(item.total or 0):,.2f}").border = border
                row += 1

            row += 1
            for label, val in [
                ("Subtotal:", f"KES {float(order.subtotal or 0):,.2f}"),
                ("Tax:", f"KES {float(order.tax or 0):,.2f}"),
                ("Discount:", f"KES {float(order.discount or 0):,.2f}"),
                ("TOTAL:", f"KES {float(order.total or 0):,.2f}"),
            ]:
                ws.cell(row=row, column=5, value=label).font = Font(bold=True)
                ws.cell(row=row, column=6, value=val)
                row += 1

            for col in range(1, 7):
                ws.column_dimensions[get_column_letter(col)].auto_size = True

            buffer = io.BytesIO()
            wb.save(buffer)
            excel_data = buffer.getvalue()
            buffer.close()

            response = current_app.response_class(
                excel_data,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response.headers["Content-Disposition"] = f"attachment; filename=INVOICE-{order.order_number}.xlsx"
            return response
        except Exception as e:
            return jsonify({"error": str(e)}), 500