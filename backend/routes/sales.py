from flask import request, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models.sales import Sale, SaleItem, SalePayment, Return, ReturnItem
from models.inventory import Product
from models.customer import Customer
from models.shop import Shop
from datetime import datetime, timezone
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import qn, OxmlElement


def utcnow():
    return datetime.now(timezone.utc)


def get_current_shop_id():
    if hasattr(current_user, "id"):
        return current_user.id
    return None


def get_shop_name():
    if hasattr(current_user, "id"):
        shop = db.session.get(Shop, current_user.id)
        if shop:
            return shop.name
    return None


def init_sales_routes(app):

    @app.route("/api/sales", methods=["GET"])
    @login_required
    def get_sales():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            status = request.args.get("status")
            start_date = request.args.get("start_date")
            end_date = request.args.get("end_date")
            customer = request.args.get("customer")

            query = Sale.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            if start_date:
                query = query.filter(Sale.created_at >= start_date)
            if end_date:
                end_datetime = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
                query = query.filter(Sale.created_at <= end_datetime)
            if customer:
                query = query.filter(Sale.customer_name.contains(customer))

            sales = query.order_by(Sale.created_at.desc()).all()
            return jsonify([s.to_dict() for s in sales])
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales/<int:sale_id>", methods=["GET"])
    @login_required
    def get_sale(sale_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            sale = Sale.query.filter_by(id=sale_id, shop_id=shop_id).first()
            if not sale:
                return jsonify({"error": "Sale not found"}), 404
            return jsonify(sale.to_dict())
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales", methods=["POST"])
    @login_required
    def create_sale():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            data = request.get_json()
            if not data or "items" not in data or len(data["items"]) == 0:
                return jsonify({"error": "At least one item is required"}), 400

            for item_data in data["items"]:
                if "product_id" not in item_data or "quantity" not in item_data:
                    return jsonify({"error": "Each item must have product_id and quantity"}), 400
                product = Product.query.filter_by(id=item_data["product_id"], shop_id=shop_id).first()
                if not product:
                    return jsonify({"error": f'Product {item_data["product_id"]} not found'}), 404
                if not product.is_active:
                    return jsonify({"error": f"Product {product.name} is inactive"}), 400
                if product.stock < item_data["quantity"]:
                    return jsonify({"error": f"Insufficient stock for {product.name}. Available: {product.stock}"}), 400

            subtotal = 0
            items = []
            for item_data in data["items"]:
                product = Product.query.filter_by(id=item_data["product_id"], shop_id=shop_id).first()
                price = float(item_data.get("price", product.price))
                quantity = item_data["quantity"]
                item_total = price * quantity
                subtotal += item_total
                items.append({"product": product, "price": price, "quantity": quantity, "total": item_total})

            discount = float(data.get("discount", 0))
            total = subtotal - discount

            customer_id = data.get("customer_id")
            customer = None
            if customer_id:
                customer = Customer.query.filter_by(id=customer_id, shop_id=shop_id).first()
                if not customer:
                    return jsonify({"error": "Customer not found"}), 404

            sale = Sale(
                shop_id=shop_id,
                customer_id=customer_id if customer else None,
                customer_name=data.get("customer_name", "Walk-in Customer"),
                customer_email=data.get("customer_email", ""),
                customer_phone=data.get("customer_phone", ""),
                payment_method=data.get("payment_method", "Cash"),
                subtotal=subtotal,
                discount=discount,
                total=total,
                notes=data.get("notes", ""),
                status="Completed",
                created_by=current_user.id if current_user.is_authenticated else None,
            )
            sale.sale_number = sale.generate_sale_number()

            db.session.add(sale)
            db.session.flush()

            for item_data in items:
                product = item_data["product"]
                db.session.add(SaleItem(
                    shop_id=shop_id,
                    sale_id=sale.id,
                    product_id=product.id,
                    product_name=product.name,
                    product_sku=product.sku,
                    quantity=item_data["quantity"],
                    price=item_data["price"],
                    cost=float(product.cost or 0),
                    total=item_data["total"],
                ))
                product.stock -= item_data["quantity"]

            db.session.add(SalePayment(
                shop_id=shop_id,
                sale_id=sale.id,
                payment_method=sale.payment_method,
                amount=total,
                status="Completed",
            ))

            if customer:
                customer_sales = Sale.query.filter_by(
                    customer_id=customer.id, shop_id=shop_id, status="Completed"
                ).all()
                customer.total_spent = sum(float(s.total or 0) for s in customer_sales)
                customer.total_orders = len(customer_sales)
                customer.update_tier()
                customer.last_activity = utcnow()

            db.session.commit()
            return jsonify(sale.to_dict()), 201
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": f"Failed to create sale: {str(e)}"}), 500

    @app.route("/api/sales/<int:sale_id>", methods=["PUT"])
    @login_required
    def update_sale(sale_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            sale = Sale.query.filter_by(id=sale_id, shop_id=shop_id).first()
            if not sale:
                return jsonify({"error": "Sale not found"}), 404

            data = request.get_json()
            for f in ["customer_name", "customer_email", "customer_phone", "payment_method", "notes"]:
                if f in data:
                    setattr(sale, f, data[f])

            if "status" in data:
                if data["status"] not in ["Completed", "Pending", "Cancelled", "Refunded"]:
                    return jsonify({"error": "Invalid status"}), 400
                sale.status = data["status"]

            db.session.commit()
            return jsonify(sale.to_dict())
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales/<int:sale_id>", methods=["DELETE"])
    @login_required
    def delete_sale(sale_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            sale = Sale.query.filter_by(id=sale_id, shop_id=shop_id).first()
            if not sale:
                return jsonify({"error": "Sale not found"}), 404
            sale.status = "Cancelled"
            db.session.commit()
            return jsonify({"message": "Sale cancelled successfully"})
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales/export/<string:format>", methods=["GET"])
    @login_required
    def export_sales(format):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            shop_name = get_shop_name() or "Sales Report"

            status = request.args.get("status")
            start_date = request.args.get("start_date")
            end_date = request.args.get("end_date")

            query = Sale.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            if start_date:
                query = query.filter(Sale.created_at >= start_date)
            if end_date:
                end_datetime = datetime.strptime(end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
                query = query.filter(Sale.created_at <= end_datetime)

            sales = query.order_by(Sale.created_at.desc()).all()
            if not sales:
                return jsonify({"error": "No sales data to export."}), 404

            if format == "excel":
                return export_excel(sales, shop_name)
            elif format == "pdf":
                return export_pdf(sales, shop_name)
            elif format == "word":
                return export_word(sales, shop_name)
            return jsonify({"error": "Invalid format"}), 400
        except Exception as e:
            return jsonify({"error": f"Failed to export: {str(e)}"}), 500

    def export_excel(sales, shop_name):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1a5276", end_color="1a5276", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(horizontal="left", vertical="center")
        number_alignment = Alignment(horizontal="right", vertical="center")
        center_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(left=Side(style="thin"), right=Side(style="thin"),
                        top=Side(style="thin"), bottom=Side(style="thin"))

        title_cell = ws.cell(row=1, column=1, value=f"SALES REPORT - {shop_name.upper()}")
        title_cell.font = Font(bold=True, size=16)
        ws.merge_cells("A1:M1")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        date_cell = ws.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        ws.merge_cells("A2:M2")
        date_cell.alignment = Alignment(horizontal="center", vertical="center")

        total_revenue = sum(float(s.total or 0) for s in sales)
        total_items = sum(len(s.items) for s in sales)
        total_qty = sum(sum(item.quantity for item in s.items) for s in sales)

        ws.cell(row=3, column=1, value=f"Total Sales: {len(sales)}")
        ws.cell(row=3, column=3, value=f"Total Revenue: KES {total_revenue:,.2f}")
        ws.cell(row=3, column=5, value=f"Total Items: {total_items}")
        ws.cell(row=3, column=7, value=f"Total Qty: {total_qty}")

        headers = ["Sale ID", "Customer", "Email", "Phone", "Date",
                   "Item Name", "SKU", "Item Qty", "Item Price", "Item Total",
                   "Payment Method", "Order Total", "Status"]

        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        row_idx = header_row + 1
        for sale in sales:
            for item in sale.items or []:
                row_data = [
                    sale.sale_number, sale.customer_name, sale.customer_email or "N/A",
                    sale.customer_phone or "N/A",
                    sale.created_at.strftime("%Y-%m-%d %H:%M") if sale.created_at else "N/A",
                    item.product_name, item.product_sku, item.quantity,
                    float(item.price or 0), float(item.price or 0) * item.quantity,
                    sale.payment_method, float(sale.total or 0), sale.status,
                ]
                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    cell.border = border
                    if isinstance(value, (int, float)):
                        cell.alignment = number_alignment
                        if col in [9, 10, 12]:
                            cell.number_format = "#,##0.00"
                    else:
                        cell.alignment = cell_alignment
                row_idx += 1

        for col in ws.columns:
            max_length = 0
            column = col[0].column
            for cell in col:
                try:
                    if isinstance(cell, openpyxl.cell.cell.MergedCell):
                        continue
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 12), 30)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f'sales_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
        )

    def export_pdf(sales, shop_name):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter),
                                rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("T", parent=styles["Heading1"], fontSize=18, alignment=1,
                                     spaceAfter=10, textColor=colors.HexColor("#1a5276"))
        elements = [Paragraph(f"{shop_name} - Sales Report", title_style)]
        date_style = ParagraphStyle("D", parent=styles["Normal"], fontSize=10, alignment=1, spaceAfter=15)
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", date_style))

        total_revenue = sum(float(s.total or 0) for s in sales)
        summary_style = ParagraphStyle("S", parent=styles["Normal"], fontSize=10, spaceAfter=15)
        elements.append(Paragraph(
            f"<b>Total Sales:</b> {len(sales)} | <b>Total Revenue:</b> KES {total_revenue:,.2f}",
            summary_style,
        ))

        table_data = [["Sale ID", "Customer", "Date", "Item", "SKU", "Qty", "Price", "Total", "Status"]]
        for sale in sales:
            for item in sale.items or []:
                table_data.append([
                    sale.sale_number, sale.customer_name[:20],
                    sale.created_at.strftime("%Y-%m-%d") if sale.created_at else "N/A",
                    item.product_name, item.product_sku, str(item.quantity),
                    f"KES {float(item.price or 0):,.2f}",
                    f"KES {float(item.price or 0) * item.quantity:,.2f}",
                    sale.status,
                ])

        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a5276")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#cccccc")),
            ("FONTSIZE", (0, 1), (-1, -1), 7),
        ]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return send_file(buffer, mimetype="application/pdf", as_attachment=True,
                         download_name=f'sales_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf')

    def export_word(sales, shop_name):
        doc = docx.Document()
        title = doc.add_heading(f"{shop_name} - Sales Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.color.rgb = RGBColor(26, 82, 118)
        doc.add_paragraph(f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}')

        total_revenue = sum(float(s.total or 0) for s in sales)
        doc.add_paragraph(f"Total Sales: {len(sales)} | Total Revenue: KES {total_revenue:,.2f}")

        table = doc.add_table(rows=1, cols=8)
        table.style = "Table Grid"
        headers = ["Sale ID", "Customer", "Date", "Item", "SKU", "Qty", "Price", "Total"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True

        for sale in sales:
            for item in sale.items or []:
                row = table.add_row().cells
                row[0].text = sale.sale_number
                row[1].text = sale.customer_name
                row[2].text = sale.created_at.strftime("%Y-%m-%d") if sale.created_at else "N/A"
                row[3].text = item.product_name
                row[4].text = item.product_sku
                row[5].text = str(item.quantity)
                row[6].text = f"KES {float(item.price or 0):,.2f}"
                row[7].text = f"KES {float(item.price or 0) * item.quantity:,.2f}"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return send_file(buffer,
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                         as_attachment=True,
                         download_name=f'sales_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx')

    @app.route("/api/returns", methods=["GET"])
    @login_required
    def get_returns():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            status = request.args.get("status")
            query = Return.query.filter_by(shop_id=shop_id)
            if status:
                query = query.filter_by(status=status)
            returns = query.order_by(Return.created_at.desc()).all()
            return jsonify([r.to_dict() for r in returns])
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/returns", methods=["POST"])
    @login_required
    def create_return():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            data = request.get_json()
            if not data or "sale_id" not in data:
                return jsonify({"error": "Sale ID is required"}), 400

            sale = Sale.query.filter_by(id=data["sale_id"], shop_id=shop_id).first()
            if not sale:
                return jsonify({"error": "Sale not found"}), 404

            if "items" not in data or len(data["items"]) == 0:
                return jsonify({"error": "At least one item required"}), 400

            total_refund = 0
            return_items = []
            for item_data in data["items"]:
                product = Product.query.filter_by(id=item_data["product_id"], shop_id=shop_id).first()
                if not product:
                    return jsonify({"error": "Product not found"}), 404
                quantity = item_data["quantity"]
                refund_amount = float(item_data.get("refund_amount", float(product.price) * quantity))
                total_refund += refund_amount
                return_items.append({
                    "product": product, "quantity": quantity,
                    "price": float(product.price), "refund_amount": refund_amount,
                    "reason": item_data.get("reason", ""),
                    "condition": item_data.get("condition", "Good"),
                })

            return_record = Return(
                shop_id=shop_id,
                sale_id=sale.id,
                customer_name=data.get("customer_name", sale.customer_name),
                customer_email=data.get("customer_email", sale.customer_email),
                customer_phone=data.get("customer_phone", sale.customer_phone),
                reason=data.get("reason", ""),
                total=total_refund,
                notes=data.get("notes", ""),
                status="Pending",
                created_by=current_user.id if current_user.is_authenticated else None,
            )
            return_record.return_number = return_record.generate_return_number()

            db.session.add(return_record)
            db.session.flush()

            for item_data in return_items:
                db.session.add(ReturnItem(
                    shop_id=shop_id,
                    return_id=return_record.id,
                    product_id=item_data["product"].id,
                    product_name=item_data["product"].name,
                    product_sku=item_data["product"].sku,
                    quantity=item_data["quantity"],
                    price=item_data["price"],
                    refund_amount=item_data["refund_amount"],
                    reason=item_data["reason"],
                    condition=item_data["condition"],
                ))
                item_data["product"].stock += item_data["quantity"]

            db.session.commit()
            return jsonify(return_record.to_dict()), 201
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": f"Failed to create return: {str(e)}"}), 500

    @app.route("/api/returns/<int:return_id>", methods=["PUT"])
    @login_required
    def update_return(return_id):
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            return_record = Return.query.filter_by(id=return_id, shop_id=shop_id).first()
            if not return_record:
                return jsonify({"error": "Return not found"}), 404

            data = request.get_json()
            if "status" in data:
                if data["status"] not in ["Pending", "Approved", "Rejected", "Completed"]:
                    return jsonify({"error": "Invalid status"}), 400
                return_record.status = data["status"]
            if "notes" in data:
                return_record.notes = data["notes"]

            db.session.commit()
            return jsonify(return_record.to_dict())
        except Exception as e:
            db.session.rollback()
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales/stats", methods=["GET"])
    @login_required
    def get_sales_stats():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401

            today = utcnow().replace(hour=0, minute=0, second=0, microsecond=0)
            today_sales = Sale.query.filter(Sale.shop_id == shop_id, Sale.created_at >= today).all()
            month_start = today.replace(day=1)
            month_sales = Sale.query.filter(Sale.shop_id == shop_id, Sale.created_at >= month_start).all()
            all_sales = Sale.query.filter_by(shop_id=shop_id).all()

            return jsonify({
                "today": {"revenue": sum(float(s.total or 0) for s in today_sales), "count": len(today_sales)},
                "month": {"revenue": sum(float(s.total or 0) for s in month_sales), "count": len(month_sales)},
                "total": {"revenue": sum(float(s.total or 0) for s in all_sales), "count": len(all_sales)},
                "status_breakdown": {
                    "completed": Sale.query.filter_by(shop_id=shop_id, status="Completed").count(),
                    "pending": Sale.query.filter_by(shop_id=shop_id, status="Pending").count(),
                    "refunded": Sale.query.filter_by(shop_id=shop_id, status="Refunded").count(),
                    "cancelled": Sale.query.filter_by(shop_id=shop_id, status="Cancelled").count(),
                },
            })
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/sales/recent", methods=["GET"])
    @login_required
    def get_recent_sales():
        try:
            shop_id = get_current_shop_id()
            if not shop_id:
                return jsonify({"error": "Shop not found"}), 401
            limit = request.args.get("limit", 10, type=int)
            sales = Sale.query.filter_by(shop_id=shop_id).order_by(Sale.created_at.desc()).limit(limit).all()
            return jsonify([s.to_dict() for s in sales])
        except Exception as e:
            return jsonify({"error": str(e)}), 500